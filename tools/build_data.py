"""สร้าง data/form2569.json + data/limits_demo.json จากไฟล์ต้นฉบับ (รันซ้ำได้)
input: ../1. 29.06.2569 แบบฟอร์มเบิกวัสดุการแพทย์ ปี 2569.xlsx, ../output/data/pcu_item_stats.csv
"""
import csv, json, math, re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WEB = Path(__file__).resolve().parents[1]
XLSX = ROOT / "1. 29.06.2569 แบบฟอร์มเบิกวัสดุการแพทย์ ปี 2569.xlsx"
STATS = ROOT / "output/data/pcu_item_stats.csv"

STEPS = [("P1", "แบบ พัสดุ 1"), ("P2", "แบบ พัสดุ 2"), ("P3", "แบบ พัสดุ 3"), ("P4", "แบบ พัสดุ 4"),
         ("P5", "แบบ พัสดุ 5"), ("CS", "แบบ จ่ายกลาง"), ("LAB", "แบบ LAB")]
PCUS = [("PCU01", "มหาดไทย", "ทั่วไป"), ("PCU02", "คลองวัว", "ทั่วไป"), ("PCU03", "บ้านแห", "ทั่วไป"),
        ("PCU04", "จำปาหล่อ", "ทั่วไป"), ("PCU05", "ย่านซื่อ", "ทั่วไป"), ("PCU06", "ศาลาแดง", "ทั่วไป"),
        ("PCU07", "ป่างิ้ว", "ทั่วไป"), ("PCU08", "โพสะ", "ทั่วไป"), ("PCU09", "หัวไผ่", "ทั่วไป"),
        ("PCU10", "บ้านยาง", "ทั่วไป"), ("PCU11", "บ้านอิฐ", "ทั่วไป"), ("PCU12", "ตลาดกรวด", "ทั่วไป"),
        ("PCU13", "บ้านรี", "ทั่วไป"), ("PCU14", "เรือนจำ", "พิเศษ"), ("PCU15", "เทศบาลเมืองฯ", "พิเศษ")]
DEMO_LIMIT_ITEMS = {"P1-01": "9", "LAB-03": "117", "LAB-04": "118"}  # item_code -> item_key(2568)


def build_form():
    wb = openpyxl.load_workbook(XLSX)
    steps, seq = [], 0
    for order, (code, sheet) in enumerate(STEPS, 1):
        ws = wb[sheet]
        rows, n = [], 0
        for r in range(12, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if a == "รวม":
                break
            if isinstance(a, (int, float)) and b:
                n += 1
                seq += 1
                assert int(a) == seq, (sheet, r, a, seq)
                rows.append({"type": "item", "code": f"{code}-{n:02d}", "seq": seq, "name": str(b),
                             "unit": ws.cell(r, 8).value, "price": float(ws.cell(r, 9).value or 0)})
            elif a is None and b:
                rows.append({"type": "section", "title": str(b)})
        steps.append({"code": code, "order": order, "sheet": sheet, "page_no": order,
                      "title": ws["A1"].value, "subject": ws["B4"].value, "to": ws["B5"].value,
                      "rows": rows})
    assert seq == 125, seq
    return {"fiscal_year": 2569, "steps": steps,
            "pcus": [{"code": c, "name": n, "print_name": n, "group": g} for c, n, g in PCUS]}


def build_limits():
    name2code = {n: c for c, n, _ in PCUS}
    out = {}
    for r in csv.DictReader(open(STATS, encoding="utf-8")):
        if r["basis"] != "2568":
            continue
        for item_code, key in DEMO_LIMIT_ITEMS.items():
            if r["item_key"] != key:
                continue
            m = math.ceil(float(r["p90_all"] or 0))
            y = math.ceil(float(r["annual_qty"] or 0))
            out.setdefault(name2code[r["pcu"]], {})[item_code] = {
                "limit_month": m or None, "limit_year": y or None,
                "note": "จำลอง: P90 รายเดือน / ยอดทั้งปี 2568"}
    return out


if __name__ == "__main__":
    (WEB / "data").mkdir(exist_ok=True)
    form = build_form()
    limits = build_limits()
    (WEB / "data/form2569.json").write_text(json.dumps(form, ensure_ascii=False, indent=1), encoding="utf-8")
    (WEB / "data/limits_demo.json").write_text(json.dumps(limits, ensure_ascii=False, indent=1), encoding="utf-8")
    print("items per step:", {s["code"]: sum(x["type"] == "item" for x in s["rows"]) for s in form["steps"]})
    print("limits:", json.dumps(limits, ensure_ascii=False))
